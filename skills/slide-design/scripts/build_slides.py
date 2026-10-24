# -*- coding: utf-8 -*-
"""slide-design 조판 스크립트

기획표(JSON) -> 16:9 pptx. 장 유형 15종(references/types.md)만 조판한다.
사용법: python3 scripts/build_slides.py plan.json out.pptx

색·글꼴은 references/theme.md 의 "설정값" 표를 읽는다. 표를 못 읽으면 아래 기본값을 쓴다.
글자 넘침 점검: 상자마다 글자 수 상한을 계산해 넘으면 경고를 출력한다(파일은 그대로 만든다).
"""
import json
import os
import re
import subprocess
import sys

# ------------------------------------------------------------ python-pptx 준비
try:
    import pptx  # noqa: F401
except ImportError:
    print("python-pptx 가 없어 설치한다: pip install python-pptx")
    cmds = [[sys.executable, "-m", "pip", "install", "-q", "python-pptx"],
            [sys.executable, "-m", "pip", "install", "-q", "--user", "python-pptx"],
            [sys.executable, "-m", "pip", "install", "-q", "--break-system-packages", "python-pptx"]]
    for c in cmds:
        if subprocess.call(c) == 0:
            break
    import site
    import importlib
    try:
        sys.path.append(site.getusersitepackages())
    except Exception:
        pass
    importlib.invalidate_caches()
    import pptx  # noqa: F401  설치 후 재시도. 여기서도 실패하면 오류로 멈춘다

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ------------------------------------------------------------ 기본값 (theme.md 가 없을 때)
FONT = "맑은 고딕"          # 기본 글꼴
INK_HEX = "#1F1D1B"         # 먹색: 제목·본문, 표지 패널·마무리 장 배경
ACC_HEX = "#C9654D"         # 강조색: 번호·키·선·강조 문장
BOX_HEX = "#F7F2EC"         # 상자색: 흐름 상자·표 강조 행·비교 머리

THEME_MD = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "references", "theme.md")


def _read_theme():
    """theme.md 의 | 이름 | 값 | 표에서 INK·ACC·BOX·FONT 를 읽는다."""
    vals = {}
    try:
        with open(THEME_MD, encoding="utf-8") as f:
            for line in f:
                m = re.match(r"^\|\s*`?(INK|ACC|BOX|FONT)`?\s*\|\s*`?([^|`]+?)`?\s*\|", line)
                if m:
                    vals[m.group(1)] = m.group(2).strip()
    except (IOError, OSError):
        pass
    return vals


def _rgb(hexstr, fallback):
    s = (hexstr or "").strip().lstrip("#")
    if not re.match(r"^[0-9A-Fa-f]{6}$", s):
        s = fallback.lstrip("#")
    return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def _mix(a, b, t):
    """색 a 와 b 를 t(0~1) 비율로 섞는다. t=0 이면 a."""
    return RGBColor(*[int(round(a[i] * (1 - t) + b[i] * t)) for i in range(3)])


_tv = _read_theme()
FONT = _tv.get("FONT", FONT) or FONT
INK = _rgb(_tv.get("INK"), INK_HEX)
ACC = _rgb(_tv.get("ACC"), ACC_HEX)
BOX = _rgb(_tv.get("BOX"), BOX_HEX)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0, 0, 0)
# 나머지 색은 세 색에서 계산한다
ACC2 = _mix(ACC, BLACK, 0.15)        # 강조 2단
MUT = _mix(INK, WHITE, 0.50)         # 보조 글자
LINE = _mix(INK, WHITE, 0.80)        # 구분선
BOX2 = _mix(BOX, INK, 0.08)          # 2단 상자
PANEL_SUB = _mix(BOX, INK, 0.15)     # 어두운 바탕 위 보조 글자

# ------------------------------------------------------------ 판형 · 여백 · 글자 크기
W, H = 13.333, 7.5
ML = MR = 0.9
CW = W - ML - MR                      # 본문 폭 11.53in
Y_LABEL, Y_TITLE, Y_BODY = 0.50, 0.80, 2.05
Y_BAND, BAND_H, Y_FOOT = 5.78, 0.66, 6.92
BODY_H = Y_BAND - 0.2 - Y_BODY        # 본문 영역 높이 3.53in

SZ_TITLE = 32      # 본문 장 제목
SZ_MSG = 20        # 핵심 문장 (띠 위)
SZ_BODY = 17       # 본문
SZ_SMALL = 12      # 라벨·하단 문구

TYPES = ("표지", "요점", "표", "흐름", "비교", "마무리",
         "구분", "질문", "인용", "수치", "카드", "타임라인", "세로 단계", "막대", "사분면")
ALIASES = {"세로단계": "세로 단계", "2x2": "사분면", "매트릭스": "사분면"}

WARN = []


# ------------------------------------------------------------ 글자 수 점검
def _units(text):
    """글자 폭 환산: 한글·전각 1, 공백 0.3, 그 밖(영문·숫자·기호) 0.55."""
    u = 0.0
    for c in str(text):
        if c == " ":
            u += 0.3
        elif ord(c) >= 0x2E80:
            u += 1.0
        else:
            u += 0.55
    return u


def _cap(width_in, size_pt, lines=1):
    """너비(인치)·글자 크기에서 들어가는 한글 글자 수."""
    per_line = width_in / (size_pt / 72.0) * 0.96
    return int(per_line * lines)


def check(where, text, limit):
    if text is None:
        return
    n = _units(text)
    if n > limit + 0.01:
        WARN.append("%s: %d자 > 상한 %d자 — \"%s\"" % (where, int(round(n)), limit, str(text)[:30]))


def _top(total_h):
    """본문 덩어리를 본문 영역 세로 가운데에 둘 때의 시작 y."""
    return Y_BODY + max(0.0, (BODY_H - total_h) / 2.0)


# ------------------------------------------------------------ 그리기 유틸
class Deck(object):
    def __init__(self, meta):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.title_line = meta.get("title_line", "")
        self.footer = meta.get("footer", "")
        self.total = 0
        self.sections = 0
        self._set_theme_fonts()

    def _set_theme_fonts(self):
        """테마 기본 글꼴을 FONT 로 바꿔 나중에 친 글자도 같은 글꼴이 되게 한다."""
        from lxml import etree
        part = self.prs.slide_masters[0].part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
        root = etree.fromstring(part.blob)
        for tag in ("majorFont", "minorFont"):
            node = root.find(".//" + qn("a:" + tag))
            if node is None:
                continue
            for child_tag in ("a:latin", "a:ea", "a:cs"):
                e = node.find(qn(child_tag))
                if e is None:
                    e = node.makeelement(qn(child_tag), {})
                    node.insert(0, e)
                e.set("typeface", FONT)
            for f in node.findall(qn("a:font")):
                if f.get("script") in ("Hang", "Hani"):
                    f.set("typeface", FONT)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def new(self):
        return self.prs.slides.add_slide(self.blank)

    def tb(self, sl, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        s = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = s.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        tf.paragraphs[0].alignment = align
        return tf

    @staticmethod
    def p(tf, first=False, before=0, after=0, line=1.3, align=None):
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        para.space_before, para.space_after, para.line_spacing = Pt(before), Pt(after), line
        if align is not None:
            para.alignment = align
        return para

    @staticmethod
    def r(para, text, size=SZ_BODY, color=INK, bold=False):
        run = para.add_run()
        run.text = str(text)
        f = run.font
        f.name, f.size, f.bold = FONT, Pt(size), bold
        f.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            e = rPr.find(qn(tag))
            if e is None:
                e = rPr.makeelement(qn(tag), {})
                rPr.append(e)
            e.set("typeface", FONT)
        return run

    def text(self, sl, x, y, w, h, txt, size, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line=1.2):
        """글자 상자 하나에 한 문단."""
        t = self.tb(sl, x, y, w, h, align=align, anchor=anchor)
        self.r(self.p(t, True, line=line, align=align), txt, size, color, bold)
        return t

    def rect(self, sl, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
        s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = fill
        s.line.fill.background()
        s.shadow.inherit = False
        st = s._element.find(qn("p:style"))   # 기본 도형 스타일(그림자·테두리 참조)을 지운다
        if st is not None:
            s._element.remove(st)
        return s

    def line(self, sl, x, y, w, color=LINE, width=0.75):
        return self.rect(sl, x, y, w, Pt(width) / 914400.0, color)

    def vline(self, sl, x, y, h, color=LINE, width=0.75):
        return self.rect(sl, x, y, Pt(width) / 914400.0, h, color)

    def badge(self, sl, x, y, d, label, size=13, fill=None):
        """번호 원."""
        e = self.rect(sl, x, y, d, d, fill or ACC, MSO_SHAPE.OVAL)
        tf = e.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        pp = tf.paragraphs[0]
        pp.alignment = PP_ALIGN.CENTER
        self.r(pp, str(label), size, WHITE, True)
        return e

    def label(self, sl, num):
        """왼쪽 위 장 번호와 제목 줄."""
        t = self.tb(sl, ML, Y_LABEL, CW, 0.3)
        para = self.p(t, True)
        self.r(para, "%02d" % num, SZ_SMALL, ACC, True)
        if self.title_line:
            self.r(para, "   " + self.title_line, SZ_SMALL, MUT)

    def chrome(self, sl, num, title, message):
        """본문 장 공통: 번호·제목 줄, 제목, 아래 핵심 문장 띠, 하단 문구, 쪽 번호."""
        self.label(sl, num)
        self.text(sl, ML, Y_TITLE, CW, 0.8, title, SZ_TITLE, INK, True, line=1.15)
        check("%02d장 제목" % num, title, _cap(CW, SZ_TITLE))
        if message:
            self.rect(sl, ML, Y_BAND, CW, BAND_H, BOX)
            self.rect(sl, ML, Y_BAND, 0.08, BAND_H, ACC)
            self.text(sl, ML + 0.35, Y_BAND, CW - 0.55, BAND_H, message, SZ_MSG, INK, True,
                      anchor=MSO_ANCHOR.MIDDLE)
            check("%02d장 핵심 문장" % num, message, _cap(CW - 0.55, SZ_MSG))
        self.foot(sl, num)

    def foot(self, sl, num=None, dark=False):
        col = PANEL_SUB if dark else MUT
        if self.footer:
            self.text(sl, ML, Y_FOOT, CW * 0.7, 0.3, self.footer, 11, col)
        if num is not None:
            self.text(sl, W - MR - 1.6, Y_FOOT, 1.6, 0.3, "%02d / %d" % (num, self.total), 11, col,
                      align=PP_ALIGN.RIGHT)

    # -------------------------------------------------------- 기존 6종
    def cover(self, num, d):
        """표지 — 왼쪽 흰 바탕 제목, 오른쪽 먹색 패널에 요점 2~3줄."""
        s = self.new()
        px = 8.6
        lw = px - ML - 0.5
        self.rect(s, px, 0, W - px, H, INK)
        self.line(s, ML, 0.95, 0.62, ACC, 3.5)
        if d.get("kicker"):
            self.text(s, ML, 1.25, lw, 0.4, d["kicker"], 15, ACC)
            check("%02d장 표지 윗줄" % num, d["kicker"], _cap(lw, 15))
        lines = _title_lines(d.get("title", ""))
        t = self.tb(s, ML, 2.1, lw, 2.6)
        for i, ln in enumerate(lines):
            self.r(self.p(t, first=(i == 0), line=1.18), ln, 44, INK, True)
            check("%02d장 표지 제목 %d줄" % (num, i + 1), ln, _cap(lw, 44))
        if len(lines) > 2:
            WARN.append("%02d장 표지 제목: %d줄 > 상한 2줄" % (num, len(lines)))
        if d.get("message"):
            self.text(s, ML, 5.25, lw, 0.9, d["message"], 16, INK)
            check("%02d장 표지 핵심 문장" % num, d["message"], _cap(lw, 16, 2))
        items = _list(d.get("items"))
        _count(num, "표지 오른쪽 항목", items, 3)
        items = items[:3]
        pw = W - px - 1.0
        n = max(len(items), 1)
        t = self.tb(s, px + 0.55, 0.6, pw, H - 1.4, anchor=MSO_ANCHOR.MIDDLE)
        for i, it in enumerate(items):
            small, big = (it[0], it[1]) if isinstance(it, (list, tuple)) and len(it) > 1 else (None, _s(it))
            if small:
                self.r(self.p(t, first=(i == 0), after=2), small, 15, PANEL_SUB)
                check("%02d장 표지 오른쪽 %d 작은 줄" % (num, i + 1), small, _cap(pw, 15))
            self.r(self.p(t, first=(i == 0 and not small), after=22 if i < n - 1 else 0, line=1.2), big, 24,
                   ACC if i == 0 else WHITE, True)
            check("%02d장 표지 오른쪽 %d" % (num, i + 1), big, _cap(pw, 24, 2))
        self.foot(s)
        return s

    def points(self, num, d):
        """요점 — 키·값 3~4행, 본문 영역 세로 가운데."""
        s = self.new()
        self.chrome(s, num, d.get("title", ""), d.get("message"))
        rows = [_kv(x) for x in _list(d.get("items"))]
        _count(num, "요점 항목", rows, 4)
        rows = rows[:5]
        n = max(len(rows), 1)
        kw = 2.7
        rh = min(0.86, BODY_H / n - 0.1)
        yy = _top(n * rh + (n - 1) * 0.1)
        for i, (k, v) in enumerate(rows):
            if k:
                self.text(s, ML, yy, kw - 0.2, rh, k, SZ_BODY - 1, ACC, True, anchor=MSO_ANCHOR.MIDDLE, line=1.15)
                check("%02d장 요점 %d 키" % (num, i + 1), k, _cap(kw - 0.2, SZ_BODY - 1, 2))
            vx = ML + (kw if k else 0)
            vw = CW - (kw if k else 0)
            self.text(s, vx, yy, vw, rh, v, SZ_BODY, INK, anchor=MSO_ANCHOR.MIDDLE)
            check("%02d장 요점 %d" % (num, i + 1), v, _cap(vw, SZ_BODY, 2 if rh >= 0.7 else 1))
            if i < n - 1:
                self.line(s, ML, yy + rh + 0.05, CW, LINE, 0.5)
            yy += rh + 0.1
        return s

    def table(self, num, d):
        """표 — 머리 1행 + 본문 2~6행. 열 폭은 글자 길이 비율로 정한다."""
        s = self.new()
        self.chrome(s, num, d.get("title", ""), d.get("message"))
        headers = [_s(h) for h in _list(d.get("headers"))]
        rows = [[_s(c) for c in _list(r)] for r in _list(d.get("rows"))]
        _count(num, "표 행", rows, 6)
        rows = rows[:7]
        nc = max([len(headers)] + [len(r) for r in rows] + [1])
        headers += [""] * (nc - len(headers))
        rows = [r + [""] * (nc - len(r)) for r in rows]
        if nc > 5:
            WARN.append("%02d장 표 열: %d개 > 상한 5개" % (num, nc))
        n = max(len(rows), 1)
        row_h = min(0.72, (BODY_H - 0.46) / n)
        size = 16 if row_h >= 0.6 else 14
        need = []
        for j in range(nc):
            m = max([_units(headers[j])] + [_units(r[j]) for r in rows] + [3.0])
            need.append(m * size / 72.0 / 0.96 + 0.35)
        if sum(need) <= CW:
            extra = (CW - sum(need)) / nc
            widths = [x + extra for x in need]
        else:
            lo, hi_w = 0.0, max(need)
            for _ in range(40):
                mid = (lo + hi_w) / 2
                if sum(min(x, mid) for x in need) > CW:
                    hi_w = mid
                else:
                    lo = mid
            widths = [min(x, lo) for x in need]
            widths[-1] += CW - sum(widths)
        hi = d.get("highlight_row")
        yy, xx = _top(0.46 + n * row_h), ML
        for j, h in enumerate(headers):
            self.text(s, xx, yy, widths[j] - 0.15, 0.3, h, 13, MUT, True)
            check("%02d장 표 머리 %d" % (num, j + 1), h, _cap(widths[j] - 0.15, 13))
            xx += widths[j]
        yy += 0.36
        self.line(s, ML, yy, CW, ACC, 1.2)
        yy += 0.1
        lines = 2 if row_h >= 0.7 else 1
        for i, row in enumerate(rows):
            on = isinstance(hi, int) and hi == i + 1
            if on:
                self.rect(s, ML - 0.12, yy + 0.02, CW + 0.24, row_h - 0.06, BOX)
            xx = ML
            for j, cell in enumerate(row):
                first = (j == 0)
                self.text(s, xx, yy, widths[j] - 0.2, row_h, cell, size, ACC2 if (on and first) else INK, first,
                          anchor=MSO_ANCHOR.MIDDLE)
                check("%02d장 표 %d행 %d열" % (num, i + 1, j + 1), cell, _cap(widths[j] - 0.2, size, lines))
                xx += widths[j]
            if i < n - 1:
                self.line(s, ML, yy + row_h, CW, LINE, 0.5)
            yy += row_h
        return s

    def flow(self, num, d):
        """흐름 — 가로 단계 3~5칸. 상자 높이는 설명 줄 수에 맞추고 세로 가운데에 둔다."""
        s = self.new()
        self.chrome(s, num, d.get("title", ""), d.get("message"))
        steps = [_step(x) for x in _list(d.get("steps") or d.get("items"))]
        _count(num, "흐름 단계", steps, 5)
        steps = steps[:6]
        n = max(len(steps), 1)
        gap = 0.42 if n <= 5 else 0.3
        bw = (CW - gap * (n - 1)) / n
        f_name = 18 if n <= 4 else 16
        f_desc = 14 if n <= 4 else 13
        maxd = max([min(len(x[1]), 3) for x in steps] + [1])
        bh = min(BODY_H, max(2.3, 1.85 + 0.42 * maxd + 0.25))
        top = _top(bh)
        for i, (name, desc) in enumerate(steps):
            x = ML + (bw + gap) * i
            last = (i == n - 1)
            self.rect(s, x, top, bw, bh, BOX2 if last else BOX)
            self.badge(s, x + 0.25, top + 0.28, 0.42, i + 1)
            iw = bw - 0.45
            self.text(s, x + 0.25, top + 0.9, iw, 0.85, name, f_name, ACC2 if last else INK, True, line=1.15)
            check("%02d장 흐름 %d단계 이름" % (num, i + 1), name, _cap(iw, f_name, 2))
            t = self.tb(s, x + 0.25, top + 1.8, iw, bh - 2.0)
            for j, ln in enumerate(desc[:3]):
                self.r(self.p(t, first=(j == 0), after=6, line=1.25), ln, f_desc, INK)
                check("%02d장 흐름 %d단계 설명 %d" % (num, i + 1, j + 1), ln, _cap(iw, f_desc, 2))
            if len(desc) > 3:
                WARN.append("%02d장 흐름 %d단계 설명: %d줄 > 상한 3줄" % (num, i + 1, len(desc)))
            if not last:
                self.text(s, x + bw, top + bh / 2 - 0.2, gap, 0.4, "▶", 16, MUT, align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE)
        return s

    def compare(self, num, d):
        """비교 — 두 열 2~4행. emphasis 로 강조할 쪽(right·left·none)을 고른다. 기본 right."""
        s = self.new()
        self.chrome(s, num, d.get("title", ""), d.get("message"))
        pairs = [_pair(x) for x in _list(d.get("pairs") or d.get("items"))]
        _count(num, "비교 행", pairs, 4)
        pairs = pairs[:5]
        emph = {"오른쪽": "right", "왼쪽": "left", "없음": "none"}.get(d.get("emphasis"), d.get("emphasis") or "right")
        has_key = any(p[0] for p in pairs)
        kw = 2.4 if has_key else 0.0
        cw2 = (CW - kw) / 2
        left, right = d.get("left", "전"), d.get("right", "후")
        n = max(len(pairs), 1)
        rh = min(0.86, (BODY_H - 0.6) / n - 0.06)
        yy = _top(0.6 + n * (rh + 0.06))
        cols = ((ML + kw, left, emph == "left", cw2 - 0.1), (ML + kw + cw2, right, emph == "right", cw2))
        for x, txt, on, w in cols:
            self.rect(s, x, yy, w, 0.46, ACC if on else BOX)
            self.text(s, x + 0.25, yy, w - 0.4, 0.46, txt, 15, WHITE if on else INK, True, anchor=MSO_ANCHOR.MIDDLE)
            check("%02d장 비교 머리" % num, txt, _cap(w - 0.4, 15))
        yy += 0.6
        for i, (k, a, b) in enumerate(pairs):
            if k:
                self.text(s, ML, yy, kw - 0.2, rh, k, 16, INK, True, anchor=MSO_ANCHOR.MIDDLE, line=1.15)
                check("%02d장 비교 %d행 항목" % (num, i + 1), k, _cap(kw - 0.2, 16, 2))
            for (x, _t, on, _w), txt, tag in ((cols[0], a, "왼쪽"), (cols[1], b, "오른쪽")):
                self.text(s, x + 0.25, yy, cw2 - 0.45, rh, txt, 17, ACC2 if on else INK, on, anchor=MSO_ANCHOR.MIDDLE)
                check("%02d장 비교 %d행 %s" % (num, i + 1, tag), txt, _cap(cw2 - 0.45, 17, 2 if rh >= 0.7 else 1))
            if i < n - 1:
                self.line(s, ML, yy + rh + 0.03, CW, LINE, 0.5)
            yy += rh + 0.06
        return s

    def closing(self, num, d):
        """마무리 — 먹색 바탕. 큰 제목, 핵심 문장, 결정·다음 할 일 2~4줄. 덩어리를 세로 가운데에 둔다."""
        s = self.new()
        self.rect(s, 0, 0, W, H, INK)
        title = d.get("title", "")
        items = [_s(x) for x in _list(d.get("items"))]
        _count(num, "마무리 항목", items, 4)
        items = items[:5]
        total = 0.3 + 1.0 + (0.8 if d.get("message") else 0) + 0.55 * len(items)
        y = max(0.9, (Y_FOOT - total) / 2)
        self.line(s, ML, y, 0.62, ACC, 3.5)
        y += 0.3
        self.text(s, ML, y, CW, 1.0, title, 44, WHITE, True, line=1.15)
        check("%02d장 마무리 제목" % num, title, _cap(CW, 44))
        y += 1.0
        if d.get("message"):
            self.text(s, ML, y, CW, 0.5, d["message"], 22, ACC, True)
            check("%02d장 마무리 핵심 문장" % num, d["message"], _cap(CW, 22))
            y += 0.8
        t = self.tb(s, ML, y, CW, 0.55 * max(len(items), 1) + 0.2)
        for i, ln in enumerate(items):
            para = self.p(t, first=(i == 0), after=10, line=1.25)
            self.r(para, "—  ", 17, PANEL_SUB)
            self.r(para, ln, 18, WHITE)
            check("%02d장 마무리 %d" % (num, i + 1), ln, _cap(CW - 0.4, 18))
        self.foot(s, dark=True)
        return s

    # -------------------------------------------------------- 신규 9종
    def section(self, num, d):
        """구분 — 상자색 바탕, 큰 덩어리 번호와 제목."""
        s = self.new()
        self.rect(s, 0, 0, W, H, BOX)
        self.sections += 1
        mark = _s(d.get("num")) or "%02d" % self.sections
        self.text(s, ML, 1.7, 4.0, 1.6, mark, 96, ACC, True, line=1.0)
        self.line(s, ML, 3.45, 0.62, ACC, 3.5)
        title = d.get("title", "")
        self.text(s, ML, 3.7, CW, 0.9, title, 40, INK, True, line=1.15)
        check("%02d장 구분 제목" % num, title, _cap(CW, 40))
        if d.get("message"):
            self.text(s, ML, 4.7, CW, 0.9, d["message"], 18, MUT, line=1.3)
            check("%02d장 구분 설명" % num, d["message"], _cap(CW, 18, 2))
        self.foot(s, num)
        return s

    def question(self, num, d):
        """질문 — 가운데 큰 질문 1~2줄과 보조 1줄."""
        s = self.new()
        self.label(s, num)
        q = d.get("question") or d.get("message") or d.get("title", "")
        sub = d.get("message") if d.get("question") else d.get("sub")
        cap = d.get("title") if d.get("question") or d.get("message") else ""
        y = 1.85
        if cap:
            self.text(s, ML, y, CW, 0.4, cap, 16, ACC, True, align=PP_ALIGN.CENTER)
            check("%02d장 질문 머리" % num, cap, _cap(CW, 16))
        self.line(s, W / 2 - 0.31, y + 0.6, 0.62, ACC, 3.5)
        self.text(s, ML + 0.5, y + 0.95, CW - 1.0, 1.9, q, 36, INK, True, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, line=1.25)
        check("%02d장 질문" % num, q, _cap(CW - 1.0, 36, 2))
        if sub:
            self.text(s, ML, y + 3.1, CW, 0.5, sub, 18, MUT, align=PP_ALIGN.CENTER)
            check("%02d장 질문 보조" % num, sub, _cap(CW, 18))
        self.foot(s, num)
        return s

    def quote(self, num, d):
        """인용 — 큰 따옴표, 인용문 1~3줄, 출처."""
        s = self.new()
        self.chrome(s, num, d.get("title", ""), d.get("message"))
        q = _s(d.get("quote") or d.get("items"))
        src = d.get("source")
        top = _top(2.9)
        self.text(s, ML, top - 0.2, 1.0, 1.0, "“", 88, ACC, True, line=1.0)
        self.text(s, ML + 1.0, top + 0.15, CW - 1.4, 1.9, q, 26, INK, True, anchor=MSO_ANCHOR.MIDDLE, line=1.3)
        check("%02d장 인용문" % num, q, _cap(CW - 1.4, 26, 3))
        if src:
            self.text(s, ML + 1.0, top + 2.3, CW - 1.4, 0.4, "— " + _s(src), 14, MUT)
            check("%02d장 인용 출처" % num, src, _cap(CW - 1.6, 14))
        return s

    def figures(self, num, d):
        """수치 — 큰 숫자 1~3개와 설명 1줄."""
        s = self.new()
        self.chrome(s, num, d.get("title", ""), d.get("message"))
        items = [_kv2(x) for x in _list(d.get("items"))]
        _count(num, "수치 항목", items, 3)
        items = items[:3]
        n = max(len(items), 1)
        cw = CW / n
        top = _top(2.2)
        for i, (val, desc) in enumerate(items):
            x = ML + cw * i
            if i > 0:
                self.vline(s, x, top + 0.1, 2.0, LINE, 0.75)
            self.text(s, x + 0.3, top, cw - 0.6, 1.2, val, 54, ACC, True, anchor=MSO_ANCHOR.BOTTOM, line=1.0)
            check("%02d장 수치 %d 값" % (num, i + 1), val, _cap(cw - 0.6, 54))
            self.text(s, x + 0.3, top + 1.35, cw - 0.6, 0.85, desc, 16, INK, line=1.25)
            check("%02d장 수치 %d 설명" % (num, i + 1), desc, _cap(cw - 0.6, 16, 2))
        return s

    def cards(self, num, d):
        """카드 — 같은 크기 상자 2~4개 나란히. 카드 이름과 설명 1~3줄."""
        s = self.new()
        self.chrome(s, num, d.get("title", ""), d.get("message"))
        items = [_step(x) for x in _list(d.get("items") or d.get("steps"))]
        _count(num, "카드", items, 4)
        items = items[:4]
        n = max(len(items), 1)
        gap = 0.3
        bw = (CW - gap * (n - 1)) / n
        maxd = max([min(len(x[1]), 3) for x in items] + [1])
        bh = min(BODY_H, max(2.0, 1.3 + 0.42 * maxd + 0.3))
        top = _top(bh)
        f_name = 20 if n <= 3 else 18
        for i, (name, desc) in enumerate(items):
            x = ML + (bw + gap) * i
            self.rect(s, x, top, bw, bh, BOX)
            self.rect(s, x, top, bw, 0.08, ACC)
            iw = bw - 0.5
            self.text(s, x + 0.25, top + 0.35, iw, 0.85, name, f_name, INK, True, line=1.15)
            check("%02d장 카드 %d 이름" % (num, i + 1), name, _cap(iw, f_name, 2))
            t = self.tb(s, x + 0.25, top + 1.3, iw, bh - 1.4)
            for j, ln in enumerate(desc[:3]):
                self.r(self.p(t, first=(j == 0), after=6, line=1.25), ln, 15, INK)
                check("%02d장 카드 %d 설명 %d" % (num, i + 1, j + 1), ln, _cap(iw, 15, 2))
            if len(desc) > 3:
                WARN.append("%02d장 카드 %d 설명: %d줄 > 상한 3줄" % (num, i + 1, len(desc)))
        return s

    def timeline(self, num, d):
        """타임라인 — 가로 선 위 날짜 점 3~6개. 위에 날짜, 아래에 사건."""
        s = self.new()
        self.chrome(s, num, d.get("title", ""), d.get("message"))
        items = [_kv2(x) for x in _list(d.get("items"))]
        _count(num, "타임라인 점", items, 6)
        items = items[:6]
        n = max(len(items), 1)
        cw = CW / n
        top = _top(2.4)
        ly = top + 0.75
        self.line(s, ML, ly, CW, LINE, 1.5)
        hi = d.get("highlight_row")
        for i, (date, ev) in enumerate(items):
            xc = ML + cw * (i + 0.5)
            on = isinstance(hi, int) and hi == i + 1
            dd = 0.26 if on else 0.2
            self.rect(s, xc - dd / 2, ly - dd / 2, dd, dd, ACC if on or not hi else MUT, MSO_SHAPE.OVAL)
            self.text(s, xc - cw / 2 + 0.1, top, cw - 0.2, 0.5, date, 16, ACC, True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.BOTTOM)
            check("%02d장 타임라인 %d 날짜" % (num, i + 1), date, _cap(cw - 0.2, 16))
            self.text(s, xc - cw / 2 + 0.1, ly + 0.3, cw - 0.2, 1.3, ev, 15, INK, on, align=PP_ALIGN.CENTER,
                      line=1.25)
            check("%02d장 타임라인 %d 사건" % (num, i + 1), ev, _cap(cw - 0.2, 15, 3))
        return s

    def vsteps(self, num, d):
        """세로 단계 — 위에서 아래로 3~4단계. 번호 · 이름 · 설명 한 행."""
        s = self.new()
        self.chrome(s, num, d.get("title", ""), d.get("message"))
        steps = [_step(x) for x in _list(d.get("steps") or d.get("items"))]
        _count(num, "세로 단계", steps, 4)
        steps = steps[:5]
        n = max(len(steps), 1)
        rh = min(0.86, BODY_H / n - 0.08)
        yy = _top(n * rh + (n - 1) * 0.08)
        nw = 3.2
        for i, (name, desc) in enumerate(steps):
            self.badge(s, ML, yy + rh / 2 - 0.22, 0.44, i + 1)
            self.text(s, ML + 0.7, yy, nw - 0.8, rh, name, 18, INK, True, anchor=MSO_ANCHOR.MIDDLE, line=1.15)
            check("%02d장 세로 단계 %d 이름" % (num, i + 1), name, _cap(nw - 0.8, 18, 2))
            dt = " ".join(desc)
            dx = ML + nw
            self.text(s, dx, yy, CW - nw, rh, dt, 16, INK, anchor=MSO_ANCHOR.MIDDLE)
            check("%02d장 세로 단계 %d 설명" % (num, i + 1), dt, _cap(CW - nw, 16, 2 if rh >= 0.7 else 1))
            if i < n - 1:
                self.line(s, ML + 0.7, yy + rh + 0.04, CW - 0.7, LINE, 0.5)
            yy += rh + 0.08
        return s

    def bars(self, num, d):
        """막대 — 항목 2~5개의 숫자 크기를 가로 막대로 비교한다."""
        s = self.new()
        self.chrome(s, num, d.get("title", ""), d.get("message"))
        items = [_bar(x, d.get("unit", "")) for x in _list(d.get("items"))]
        _count(num, "막대", items, 5)
        items = items[:6]
        n = max(len(items), 1)
        rh = min(0.7, BODY_H / n - 0.12)
        yy = _top(n * rh + (n - 1) * 0.12)
        lw, vw = 2.6, 1.7
        bmax = CW - lw - vw
        top_v = max([abs(v) for _, v, _ in items] + [1e-9])
        hi = d.get("highlight_row")
        for i, (name, v, shown) in enumerate(items):
            on = (isinstance(hi, int) and hi == i + 1) or (not hi)
            self.text(s, ML, yy, lw - 0.2, rh, name, 16, INK, True, anchor=MSO_ANCHOR.MIDDLE, line=1.15)
            check("%02d장 막대 %d 항목" % (num, i + 1), name, _cap(lw - 0.2, 16, 2))
            bw = max(0.06, bmax * abs(v) / top_v)
            bh = rh * 0.62
            self.rect(s, ML + lw, yy + (rh - bh) / 2, bw, bh, ACC if on else BOX2)
            self.text(s, ML + lw + bw + 0.15, yy, vw, rh, shown, 17, ACC2 if on else INK, True,
                      anchor=MSO_ANCHOR.MIDDLE)
            check("%02d장 막대 %d 값" % (num, i + 1), shown, _cap(vw, 17))
            yy += rh + 0.12
        if any(v is None for _, v, _ in items):
            WARN.append("%02d장 막대: 숫자가 아닌 값이 있다" % num)
        return s

    def quadrant(self, num, d):
        """사분면 — 2×2 칸. cells 순서: 왼쪽 위, 오른쪽 위, 왼쪽 아래, 오른쪽 아래."""
        s = self.new()
        self.chrome(s, num, d.get("title", ""), d.get("message"))
        cells = [_kv2(x) for x in _list(d.get("cells") or d.get("items"))]
        if len(cells) != 4:
            WARN.append("%02d장 사분면 칸: %d개 (4개여야 한다)" % (num, len(cells)))
        cells = (cells + [("", "")] * 4)[:4]
        ax, ay = _s(d.get("x_axis")), _s(d.get("y_axis"))
        hi = d.get("highlight")
        gx = ML + (0.45 if ay else 0)
        gw = CW - (0.45 if ay else 0)
        gh = BODY_H - (0.4 if ax else 0)
        top = Y_BODY
        cw, ch = (gw - 0.12) / 2, (gh - 0.12) / 2
        for i, (name, desc) in enumerate(cells):
            x = gx + (cw + 0.12) * (i % 2)
            y = top + (ch + 0.12) * (i // 2)
            on = isinstance(hi, int) and hi == i + 1
            self.rect(s, x, y, cw, ch, ACC if on else BOX)
            self.text(s, x + 0.25, y + 0.15, cw - 0.5, 0.45, name, 18, WHITE if on else INK, True)
            check("%02d장 사분면 %d 이름" % (num, i + 1), name, _cap(cw - 0.5, 18))
            self.text(s, x + 0.25, y + 0.62, cw - 0.5, ch - 0.7, desc, 15, WHITE if on else INK, line=1.25)
            check("%02d장 사분면 %d 설명" % (num, i + 1), desc, _cap(cw - 0.5, 15, 2))
        if ax:
            self.text(s, gx, top + gh + 0.08, gw, 0.3, "가로: " + ax + "  →", 13, MUT, True, align=PP_ALIGN.RIGHT)
        if ay:
            self.text(s, ML - 1.0, top + gh / 2 - 0.15, 2.3, 0.3, "세로: " + ay + "  →", 13, MUT, True,
                      align=PP_ALIGN.CENTER)
            s.shapes[-1].rotation = 270.0
        return s


# ------------------------------------------------------------ 입력 정리
def _s(x):
    if isinstance(x, (list, tuple)):
        return " ".join(_s(i) for i in x)
    if isinstance(x, dict):
        return " ".join(_s(v) for v in x.values())
    return "" if x is None else str(x)


def _list(x):
    if x is None:
        return []
    return list(x) if isinstance(x, (list, tuple)) else [x]


def _kv(x):
    """요점 항목: "키: 값" 문자열, [키, 값], {"key","value"} 모두 받는다."""
    if isinstance(x, dict):
        return _s(x.get("key", "")), _s(x.get("value", x.get("text", "")))
    if isinstance(x, (list, tuple)):
        return (_s(x[0]), _s(x[1])) if len(x) > 1 else ("", _s(x[0]))
    m = re.match(r"^\s*([^:：]{1,12})\s*[:：]\s*(.+)$", _s(x))
    return (m.group(1), m.group(2)) if m else ("", _s(x))


def _step(x):
    """흐름 단계: {"name","desc"}, [이름, 설명...], "이름: 설명" 모두 받는다."""
    if isinstance(x, dict):
        desc = x.get("desc", x.get("items", []))
        return _s(x.get("name", x.get("title", ""))), [_s(i) for i in _list(desc)]
    if isinstance(x, (list, tuple)):
        rest = list(x[1:])
        if len(rest) == 1 and isinstance(rest[0], (list, tuple)):
            rest = list(rest[0])
        return _s(x[0]), [_s(i) for i in rest]
    k, v = _kv(x)
    return (k, [v]) if k else (v, [])


def _pair(x):
    """비교 행: [항목, 왼쪽, 오른쪽] 또는 [왼쪽, 오른쪽] 또는 {"item","left","right"}."""
    if isinstance(x, dict):
        return _s(x.get("item", "")), _s(x.get("left", "")), _s(x.get("right", ""))
    x = _list(x)
    if len(x) >= 3:
        return _s(x[0]), _s(x[1]), _s(x[2])
    if len(x) == 2:
        return "", _s(x[0]), _s(x[1])
    return "", _s(x[0]) if x else "", ""


def _title_lines(title):
    """표지 제목을 2줄로 나눈다. 줄바꿈(\\n)이 있으면 그대로, 없으면 가운데에 가까운 띄어쓰기에서 나눈다."""
    title = _s(title)
    if "\n" in title:
        return [t.strip() for t in title.split("\n") if t.strip()]
    if _units(title) <= 10:
        return [title]
    words = title.split(" ")
    best, bi = None, 1
    for i in range(1, len(words)):
        a, b = " ".join(words[:i]), " ".join(words[i:])
        diff = max(_units(a), _units(b))
        if best is None or diff < best:
            best, bi = diff, i
    return [" ".join(words[:bi]), " ".join(words[bi:])] if len(words) > 1 else [title]


def _count(num, what, seq, limit):
    if len(seq) > limit:
        WARN.append("%02d장 %s: %d개 > 상한 %d개" % (num, what, len(seq), limit))


def _kv2(x):
    """두 칸 항목(값·설명, 날짜·사건, 칸 이름·설명): [a, b], {"a","b"} 또는 "a: b"."""
    if isinstance(x, dict):
        vals = list(x.values())
        return (_s(vals[0]), _s(vals[1])) if len(vals) > 1 else (_s(vals[0]) if vals else "", "")
    if isinstance(x, (list, tuple)):
        return (_s(x[0]), _s(x[1:])) if len(x) > 1 else (_s(x[0]), "")
    k, v = _kv(x)
    return (k, v) if k else (v, "")


def _bar(x, unit=""):
    """막대 항목: [이름, 숫자] 또는 [이름, 숫자, 표시 문구]. 표시 문구가 없으면 숫자+단위."""
    x = _list(x) if not isinstance(x, dict) else [x.get("name", ""), x.get("value"), x.get("label")]
    name = _s(x[0]) if x else ""
    raw = x[1] if len(x) > 1 else None
    try:
        v = float(str(raw).replace(",", ""))
    except (TypeError, ValueError):
        v = 0.0
    shown = _s(x[2]) if len(x) > 2 and x[2] not in (None, "") else (_s(raw) + unit)
    return name, v, shown


# ------------------------------------------------------------ 실행
def build(plan, out):
    meta = plan.get("meta", {}) if isinstance(plan, dict) else {}
    slides = plan.get("slides", []) if isinstance(plan, dict) else plan
    if not slides:
        sys.exit("오류: 기획표에 장(slides)이 없다")
    for sd in slides:
        sd["type"] = ALIASES.get(sd.get("type"), sd.get("type"))
    bad = [(i + 1, sd.get("type")) for i, sd in enumerate(slides) if sd.get("type") not in TYPES]
    if bad:
        sys.exit("오류: 장 유형은 %s 중 하나여야 한다. 맞지 않는 장: %s" % ("·".join(TYPES), bad))
    if not 5 <= len(slides) <= 8:
        WARN.append("장 수: %d장 (기준 5~8장)" % len(slides))
    deck = Deck(meta)
    deck.total = len(slides)
    fn = {"표지": deck.cover, "요점": deck.points, "표": deck.table,
          "흐름": deck.flow, "비교": deck.compare, "마무리": deck.closing,
          "구분": deck.section, "질문": deck.question, "인용": deck.quote, "수치": deck.figures,
          "카드": deck.cards, "타임라인": deck.timeline, "세로 단계": deck.vsteps, "막대": deck.bars,
          "사분면": deck.quadrant}
    for i, sd in enumerate(slides):
        fn[sd["type"]](i + 1, sd)
        notes = _s(sd.get("notes"))
        if notes:
            deck.prs.slides[-1].notes_slide.notes_text_frame.text = notes
        if i >= 2 and slides[i - 2]["type"] == slides[i - 1]["type"] == sd["type"]:
            WARN.append("%02d장 장 유형: %s 3장 연속" % (i + 1, sd["type"]))
    deck.prs.save(out)
    print("생성: %s (%d장, 글꼴 %s)" % (out, len(slides), FONT))
    if WARN:
        print("넘침 경고 %d건" % len(WARN))
        for w in WARN:
            print("- " + w)
    else:
        print("넘침 경고 0건")


def main():
    if len(sys.argv) != 3:
        sys.exit("사용법: python3 scripts/build_slides.py plan.json out.pptx")
    with open(sys.argv[1], encoding="utf-8") as f:
        plan = json.load(f)
    build(plan, sys.argv[2])


if __name__ == "__main__":
    main()
